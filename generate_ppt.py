#!/usr/bin/env python3
"""
Automated PowerPoint Presentation Builder for Project 2:
Generates docs/REPORT_Presentation.pptx with a dark cybersecurity theme
and embeds terminal screenshots from docs/screenshots/.
"""

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# Dark Cybersecurity Theme Colors
BG_DARK = RGBColor(15, 23, 42)       # Slate 900
ACCENT_CYAN = RGBColor(56, 189, 248) # Sky 400
ACCENT_GREEN = RGBColor(74, 222, 128)# Green 400
ACCENT_RED = RGBColor(248, 113, 113) # Red 400
TEXT_WHITE = RGBColor(241, 245, 249) # Slate 100
TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400
PANEL_BG = RGBColor(30, 41, 59)      # Slate 800


def apply_dark_background(slide) -> None:
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK


def add_header(slide, title_text: str, subtitle_text: str = "") -> None:
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.9))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p.font.name = "Consolas"

    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(14)
        p2.font.color.rgb = TEXT_MUTED
        p2.font.name = "Calibri"


def add_bullet_card(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    heading: str,
    bullets: list[str],
    heading_color: RGBColor = ACCENT_GREEN,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True

    p_head = tf.paragraphs[0]
    p_head.text = heading
    p_head.font.size = Pt(18)
    p_head.font.bold = True
    p_head.font.color.rgb = heading_color
    p_head.font.name = "Consolas"
    p_head.space_after = Pt(8)

    for b_text in bullets:
        p = tf.add_paragraph()
        p.text = f"• {b_text}"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.font.name = "Calibri"
        p.space_after = Pt(6)


def add_screenshot_if_exists(
    slide, img_path: Path, left: float, top: float, width: float
) -> None:
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(left), Inches(top), width=Inches(width))
    else:
        placeholder = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(3.5))
        tf = placeholder.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"[Screenshot Placeholder: {img_path.name}]\nPlace image in docs/screenshots/ and re-run generate_ppt.py"
        p.font.size = Pt(14)
        p.font.color.rgb = ACCENT_RED


def build_presentation() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    shots_dir = Path("docs/screenshots")

    # ---------------------------------------------------------
    # SLIDE 1: Title Slide
    # ---------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    apply_dark_background(s1)
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.5))
    tf = t_box.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "PROJECT 2: SECURE MESSAGING WITH DIGITAL SIGNATURES"
    p0.font.size = Pt(34)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN
    p0.font.name = "Consolas"
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "Ed25519 (RFC 8032) Authenticity, AES-256-GCM Encrypt-then-Sign & Active MITM Attack Evaluation"
    p1.font.size = Pt(20)
    p1.font.color.rgb = TEXT_WHITE
    p1.font.name = "Calibri"
    p1.space_after = Pt(24)

    p2 = tf.add_paragraph()
    p2.text = "Stack: Python 3.10+ | cryptography (OpenSSL Hazmat) | Rich CLI Sandbox"
    p2.font.size = Pt(15)
    p2.font.color.rgb = ACCENT_GREEN
    p2.font.name = "Consolas"

    # ---------------------------------------------------------
    # SLIDE 2: System Design & Why Ed25519?
    # ---------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    apply_dark_background(s2)
    add_header(s2, "1. System Architecture & Primitive Selection", "Why Ed25519 (RFC 8032) + AES-256-GCM over Legacy RSA / ECDSA")
    add_bullet_card(
        s2, 0.6, 1.5, 5.8, 5.2,
        "Why Ed25519 Digital Signatures?",
        [
            "128-Bit Security Level: Equivalent to 3072-bit RSA with compact 32-byte public keys and 64-byte signatures.",
            "Deterministic Nonce (No k-Reuse Bug): Derives r = SHA-512(prefix || M) mod L, eliminating RNG failures like the Sony PS3 ECDSA breach.",
            "Side-Channel Resistance: Twisted Edwards complete addition formulas prevent branch/timing leaks.",
            "Verification Equation: Checks [S]B = R + [k]A where k = SHA-512(R || A || M) mod L.",
        ],
        ACCENT_GREEN,
    )
    add_bullet_card(
        s2, 6.8, 1.5, 5.9, 5.2,
        "3-Layer Modular Codebase",
        [
            "crypto_core.py: Ed25519 KeyGen/Sign/Verify, AES-256-GCM Encrypt-then-Sign, SHA-256 Fingerprint Directory, Replay Cache.",
            "experiments.py: Automated Dolev-Yao MITM attack testbed covering Baseline + 5 security experiments & JSON/PUB artifact exporter.",
            "app.py: Interactive Rich CLI simulator with live Attacker Sandbox (Tamper, Bit-Flip, Impersonate, Key-Swap, Replay).",
            "Strict Key Hygiene: Ephemeral RAM private keys; .gitignore blocks *.pem while tracking *.pub.",
        ],
        ACCENT_CYAN,
    )

    # ---------------------------------------------------------
    # SLIDE 3: Envelope Specification & Canonicalization
    # ---------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    apply_dark_background(s3)
    add_header(s3, "2. Wire Envelope Schema & Canonical Serialization", "Binding Metadata Inside the Cryptographic Boundary")
    add_bullet_card(
        s3, 0.6, 1.5, 6.0, 5.2,
        "Deterministic JSON Canonicalization",
        [
            "Ed25519 signs raw byte arrays, not abstract dictionaries.",
            "Rule: json.dumps(envelope, sort_keys=True, separators=(',', ':')).encode('utf-8')",
            "Lexicographically sorts keys and strips whitespace so Alice and Bob hash the exact same byte sequence.",
            "Strict Verification Order: (1) Packet Schema -> (2) SHA-256 Key Fingerprint -> (3) Ed25519 Signature -> (4) Replay Counter/Timestamp -> (5) AES-GCM Decrypt.",
        ],
        ACCENT_GREEN,
    )
    add_bullet_card(
        s3, 6.8, 1.5, 5.9, 5.2,
        "Why Metadata MUST Be Inside the Signed Envelope",
        [
            "sender & recipient: Prevents Surreptitious Forwarding / Unknown Key-Share routing tampering.",
            "counter, timestamp, msg_id: Prevents an attacker from incrementing the counter on a captured packet to bypass replay detection.",
            "encrypted & nonce: Prevents downgrade attacks or GCM IV manipulation prior to signature verification.",
        ],
        ACCENT_RED,
    )

    # ---------------------------------------------------------
    # SLIDE 4: Baseline Authentic Exchange (exp0)
    # ---------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    apply_dark_background(s4)
    add_header(s4, "3. Baseline Test: Authentic Signed Exchange (exp0)", "Normal Alice -> Bob Signing & Verification Flow")
    add_bullet_card(
        s4, 0.6, 1.5, 5.2, 5.2,
        "Baseline Verification Pipeline",
        [
            "Hypothesis: An unmodified packet signed by Alice's Ed25519 private key passes all checks.",
            "1. Bob fetches Alice's public key and verifies SHA-256(PK_Alice) against trusted out-of-band fingerprint.",
            "2. Reconstructs canonical envelope bytes and verifies 64-byte Ed25519 signature.",
            "3. Checks counter > last_counter and |now - timestamp| <= 300s.",
            "Result: valid=True, status_code='SUCCESS'.",
        ],
        ACCENT_GREEN,
    )
    add_screenshot_if_exists(s4, shots_dir / "01_normal_exchange.png", 6.0, 1.5, 6.7)

    # ---------------------------------------------------------
    # SLIDE 5: Exp 1A — Message Payload Modification
    # ---------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    apply_dark_background(s5)
    add_header(s5, "4. Experiment 1A: Payload Tampering in Transit", "MITM Modifies Message Text Without Re-Signing")
    add_bullet_card(
        s5, 0.6, 1.5, 5.2, 5.2,
        "Attack & Cryptographic Proof",
        [
            "Attack: Mallory intercepts Alice's packet and alters payload ('Auth payment of $99999 to Mallory') without Alice's private key.",
            "Observed Result: REJECTED (status_code='INVALID_SIGNATURE', plaintext=None).",
            "Math Explanation: Modifying 1 byte changes canonical M -> M'. By the SHA-512 avalanche effect, ~50% of bits flip in k' = SHA-512(R || A || M').",
            "Thus [S]B != R + [k']A on Curve25519.",
        ],
        ACCENT_RED,
    )
    add_screenshot_if_exists(s5, shots_dir / "02_tampering_rejected.png", 6.0, 1.5, 6.7)

    # ---------------------------------------------------------
    # SLIDE 6: Exp 1B — Signature Byte Corruption
    # ---------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    apply_dark_background(s6)
    add_header(s6, "5. Experiment 1B: Signature Bit-Flip Corruption", "MITM Flips 1 Byte (XOR 0xFF) in the 64-Byte Ed25519 Signature")
    add_bullet_card(
        s6, 0.6, 1.5, 5.2, 5.2,
        "Attack & Cryptographic Proof",
        [
            "Attack: Payload remains untouched, but Mallory flips byte 0 (XOR 0xFF) of the 64-byte signature.",
            "Observed Result: REJECTED (status_code='INVALID_SIGNATURE').",
            "Math Explanation: The 64-byte signature encodes curve point R (32B) and scalar S (32B).",
            "Corrupting byte 0 mutates R -> R', breaking the verification equality [S]B = R' + [k]A.",
        ],
        ACCENT_RED,
    )
    add_screenshot_if_exists(s6, shots_dir / "02b_corrupted_signature_rejected.png", 6.0, 1.5, 6.7)

    # ---------------------------------------------------------
    # SLIDE 7: Exp 2 — Sender Impersonation Attempt
    # ---------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    apply_dark_background(s7)
    add_header(s7, "6. Experiment 2: Sender Impersonation Attempt", "Mallory Signs Envelope with SK_Mallory While Claiming 'sender': 'Alice'")
    add_bullet_card(
        s7, 0.6, 1.5, 5.2, 5.2,
        "Attack & Cryptographic Proof",
        [
            "Attack: Mallory sets envelope['sender'] = 'Alice' ('Send production SSH pvt key...') and signs with Mallory's own private key.",
            "Observed Result: REJECTED (status_code='INVALID_SIGNATURE').",
            "Math Explanation: Ed25519 is EUF-CMA secure under the Elliptic Curve Discrete Logarithm Problem (ECDLP).",
            "Mallory's signature satisfies [S]B = R + [k]A_Mallory, which fails when Bob verifies against A_Alice.",
        ],
        ACCENT_RED,
    )
    add_screenshot_if_exists(s7, shots_dir / "03_impersonation_rejected.png", 6.0, 1.5, 6.7)

    # ---------------------------------------------------------
    # SLIDE 8: Exp 3 — Public-Key Substitution & Fingerprint Defense
    # ---------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    apply_dark_background(s8)
    add_header(s8, "7. Experiment 3: Public-Key Substitution Attack", "Directory Poisoning vs. Out-of-Band SHA-256 Fingerprint Pinning")
    add_bullet_card(
        s8, 0.6, 1.5, 5.2, 5.2,
        "The Trust Bootstrap Problem",
        [
            "Attack: Mallory overwrites Alice's public key in PublicKeyDirectory with PK_Mallory and signs wire transfer command.",
            "Part A (verify_key_fingerprint=False): Bob blindly trusts directory -> ACCEPTED (VULNERABLE!).",
            "Part B (verify_key_fingerprint=True): Bob compares SHA-256(PK_Directory) against pinned out-of-band fingerprint.",
            "Result: REJECTED (status_code='KEY_SUBSTITUTION_DETECTED').",
        ],
        ACCENT_RED,
    )
    add_screenshot_if_exists(s8, shots_dir / "04_key_substitution.png", 6.0, 1.5, 6.7)

    # ---------------------------------------------------------
    # SLIDE 9: Exp 4 — Confidentiality & AES-256-GCM Extension
    # ---------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    apply_dark_background(s9)
    add_header(s9, "8. Experiment 4: Confidentiality & AES-256-GCM Extension", "Why Signatures Alone Expose Plaintext & How Encrypt-then-Sign Fixes It")
    add_bullet_card(
        s9, 0.6, 1.5, 5.2, 5.2,
        "Encrypt-then-Sign (EtS) Architecture",
        [
            "Part A (Signature Only): Eavesdropper reads 'payload' in cleartext JSON on the wire.",
            "Part B (AES-256-GCM + Ed25519): Plaintext encrypted with 256-bit key & 96-bit random nonce before signing.",
            "Wire shows only Base64 ciphertext + nonce (IND-CPA secure).",
            "Bob verifies outer Ed25519 signature FIRST, then decrypts AES-GCM payload -> VALID (SUCCESS).",
        ],
        ACCENT_GREEN,
    )
    add_screenshot_if_exists(s9, shots_dir / "05_confidentiality_aesgcm.png", 6.0, 1.5, 6.7)

    # ---------------------------------------------------------
    # SLIDE 10: Exp 5 — Replay Attack & Counter Defense
    # ---------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    apply_dark_background(s10)
    add_header(s10, "9. Experiment 5: Replay Attack & Counter Defense", "Stateful Sequence Tracking & Signed Counter Binding")
    add_bullet_card(
        s10, 0.6, 1.5, 5.2, 5.2,
        "3-Part Replay Evaluation",
        [
            "Attack: Mallory captures a valid signed packet (counter=2) and retransmits it to Bob.",
            "Part A (Defense OFF): Accepted a second time (VULNERABLE!).",
            "Part B (Defense ON): Bob checks seen_messages & counter <= last_counter -> REJECTED ('REPLAY_DETECTED').",
            "Part C (Counter Tampering): Mallory increments counter += 1 without SK_Alice -> REJECTED ('INVALID_SIGNATURE').",
        ],
        ACCENT_RED,
    )
    add_screenshot_if_exists(s10, shots_dir / "06_replay_defense.png", 6.0, 1.5, 6.7)

    # ---------------------------------------------------------
    # SLIDE 11: Deep-Dive Security Questions
    # ---------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    apply_dark_background(s11)
    add_header(s11, "10. Core Cryptographic Design Takeaways", "Key Answers to Syllabus Security Questions")
    add_bullet_card(
        s11, 0.6, 1.5, 5.9, 5.3,
        "1. Authenticity vs. Confidentiality & Trust Models",
        [
            "Signatures != Encryption: Verify(PK, M, sig) requires M in cleartext; AES-256-GCM provides confidentiality while Ed25519 provides non-repudiation.",
            "Why EtS Order? Signing the ciphertext prevents attackers from feeding unauthenticated ciphertexts into the AES-GCM decryption routine.",
            "Trust Scaling: OOB Fingerprints (high security, O(N^2) manual effort) -> X.509 PKI CAs (O(N) scalable, single point of compromise) -> Key Transparency (append-only Merkle Tree logs).",
        ],
        ACCENT_CYAN,
    )
    add_bullet_card(
        s11, 6.8, 1.5, 5.9, 5.3,
        "2. Replay Defense State & Sync Trade-Offs",
        [
            "Monotonic Counters: O(1) memory per sender, but sensitive to out-of-order network delivery unless paired with a sliding window bitmap.",
            "UUIDv4 Nonces (msg_id): Tolerates out-of-order packets effortlessly, but requires O(M) cache memory growth unless bounded by TTL.",
            "UTC Timestamps (300s Window): Bounds nonce cache memory, but requires secure NTP clock synchronization to avoid clock-skew DoS.",
        ],
        ACCENT_GREEN,
    )

    output_path = Path("docs/REPORT_Presentation.pptx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"[+] Successfully generated 11-slide deck at: {output_path}")


if __name__ == "__main__":
    build_presentation()